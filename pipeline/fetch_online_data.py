#!/usr/bin/env python3
"""Fetch everything AVONET does not carry, for every bird in the workbook.

Produces pipeline/raw/online_data.json: one record per scientific name with
the English common name, its Wikipedia article, IUCN status, 12-month pageview
total and the continent its range centroid falls in.

This is the only step that touches the network, and it is slow -- tens of
thousands of requests. It is meant to run once, in CI (GitHub Actions runners
have open internet), and commit its output so every later build is offline and
reproducible.

    python3 pipeline/fetch_online_data.py            # full run
    python3 pipeline/fetch_online_data.py --limit 50 # smoke test

Progress is checkpointed to pipeline/.cache/, so an interrupted run resumes
instead of starting over.

Sources: Wikidata (CC0) for names, articles and IUCN status; the Wikimedia
pageviews API for popularity; Natural Earth for continent boundaries.
"""

from __future__ import annotations

import argparse
import json
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, Iterable

REPO = Path(__file__).resolve().parent.parent
RAW = REPO / "pipeline" / "raw"
CACHE = REPO / "pipeline" / ".cache"
WORKBOOK = RAW / "AVONET Supplementary dataset 1.xlsx"
OUTPUT = RAW / "online_data.json"

USER_AGENT = (
    "daily-bird-pipeline/1.0 (https://github.com/julesfll/daily-bird) "
    "python-urllib"
)
WIKIDATA_SPARQL = "https://query.wikidata.org/sparql"
PAGEVIEWS = (
    "https://wikimedia.org/api/rest_v1/metrics/pageviews/per-article"
    "/en.wikipedia/all-access/all-agents/{title}/monthly/{start}/{end}"
)
# Natural Earth at 50m keeps small islands that 110m drops -- which matters
# when a fifth of the pool are island endemics and seabirds.
NATURAL_EARTH = [
    "https://naturalearth.s3.amazonaws.com/50m_cultural/ne_50m_admin_0_countries.zip",
    "https://naturalearth.s3.amazonaws.com/110m_cultural/ne_110m_admin_0_countries.zip",
]

PAGEVIEW_START, PAGEVIEW_END = "2025010100", "2025123100"
SPARQL_BATCH = 120
PAGEVIEW_WORKERS = 8


def http_get(url: str, *, timeout: int = 120, retries: int = 4) -> bytes:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    delay = 2.0
    for attempt in range(retries):
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return response.read()
        except urllib.error.HTTPError as exc:
            # 404 means the article simply does not exist; retrying won't help.
            if exc.code == 404:
                raise
            if attempt == retries - 1:
                raise
        except Exception:
            if attempt == retries - 1:
                raise
        time.sleep(delay)
        delay *= 2
    raise RuntimeError("unreachable")


# ---------------------------------------------------------------- AVONET


def to_float(value: Any) -> float | None:
    """AVONET writes missing numbers as the string 'NA', not as a blank."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def load_species() -> list[dict[str, Any]]:
    import openpyxl

    wb = openpyxl.load_workbook(WORKBOOK, read_only=True)
    ws = wb["AVONET1_BirdLife"]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    at = {name: i for i, name in enumerate(header)}

    out = []
    missing = 0
    for row in rows:
        name = row[at["Species1"]]
        if not name:
            continue
        lat = to_float(row[at["Centroid.Latitude"]])
        lon = to_float(row[at["Centroid.Longitude"]])
        if lat is None or lon is None:
            missing += 1
        out.append({"sci": name, "lat": lat, "lon": lon})
    if missing:
        print(f"  {missing} species have no usable range centroid")
    return out


# ---------------------------------------------------------------- Wikidata

SPARQL = """
SELECT ?taxonName ?vernacular ?article ?iucnLabel WHERE {
  VALUES ?taxonName { %s }
  ?item wdt:P225 ?taxonName .
  OPTIONAL { ?item wdt:P1843 ?vernacular . FILTER(LANG(?vernacular) = "en") }
  OPTIONAL {
    ?article schema:about ?item ;
             schema:isPartOf <https://en.wikipedia.org/> .
  }
  OPTIONAL { ?item wdt:P141 ?iucn . }
  SERVICE wikibase:label { bd:serviceParam wikibase:language "en". }
}
"""


def fetch_wikidata(names: list[str]) -> dict[str, dict[str, str]]:
    CACHE.mkdir(parents=True, exist_ok=True)
    found: dict[str, dict[str, str]] = {}

    batches = [names[i : i + SPARQL_BATCH] for i in range(0, len(names), SPARQL_BATCH)]
    for index, batch in enumerate(batches, 1):
        cached = CACHE / f"wikidata-{index:04d}.json"
        if cached.exists():
            payload = json.loads(cached.read_text(encoding="utf-8"))
        else:
            values = " ".join(json.dumps(n) for n in batch)
            query = SPARQL % values
            url = f"{WIKIDATA_SPARQL}?{urllib.parse.urlencode({'query': query, 'format': 'json'})}"
            try:
                payload = json.loads(http_get(url))
            except Exception as exc:  # keep going; a lost batch is not fatal
                print(f"  wikidata batch {index}/{len(batches)} failed: {exc}", file=sys.stderr)
                continue
            cached.write_text(json.dumps(payload), encoding="utf-8")
            time.sleep(1.0)  # a free public endpoint; don't hammer it

        for binding in payload["results"]["bindings"]:
            name = binding["taxonName"]["value"]
            record = found.setdefault(name, {})
            if "vernacular" in binding and "name" not in record:
                record["name"] = binding["vernacular"]["value"]
            if "article" in binding and "wiki" not in record:
                title = binding["article"]["value"].rsplit("/", 1)[-1]
                record["wiki"] = urllib.parse.unquote(title).replace("_", " ")
            if "iucnLabel" in binding and "status" not in record:
                record["status"] = binding["iucnLabel"]["value"]

        if index % 10 == 0 or index == len(batches):
            print(f"  wikidata {index}/{len(batches)} batches, {len(found)} matched")

    return found


# ---------------------------------------------------------------- pageviews


def fetch_pageviews(titles: list[str]) -> dict[str, int]:
    store = CACHE / "pageviews.json"
    views: dict[str, int] = {}
    if store.exists():
        views = json.loads(store.read_text(encoding="utf-8"))

    todo = [t for t in titles if t not in views]
    print(f"  pageviews: {len(views)} cached, {len(todo)} to fetch")

    def one(title: str) -> tuple[str, int]:
        quoted = urllib.parse.quote(title.replace(" ", "_"), safe="")
        url = PAGEVIEWS.format(title=quoted, start=PAGEVIEW_START, end=PAGEVIEW_END)
        try:
            payload = json.loads(http_get(url, timeout=60, retries=2))
            return title, sum(item["views"] for item in payload.get("items", []))
        except Exception:
            return title, 0  # missing or renamed article simply ranks last

    done = 0
    with ThreadPoolExecutor(max_workers=PAGEVIEW_WORKERS) as pool:
        for title, total in pool.map(one, todo):
            views[title] = total
            done += 1
            if done % 500 == 0:
                print(f"  pageviews {done}/{len(todo)}")
                store.write_text(json.dumps(views), encoding="utf-8")

    store.write_text(json.dumps(views), encoding="utf-8")
    return views


# ---------------------------------------------------------------- continent


def build_continents():
    from shapely.geometry import shape
    from shapely.ops import unary_union
    import shapefile  # pyshp

    CACHE.mkdir(parents=True, exist_ok=True)
    for url in NATURAL_EARTH:
        stem = url.rsplit("/", 1)[-1].replace(".zip", "")
        archive = CACHE / f"{stem}.zip"
        try:
            if not archive.exists():
                archive.write_bytes(http_get(url, timeout=180))
            zf = zipfile.ZipFile(archive)
            reader = shapefile.Reader(
                shp=zf.open(f"{stem}.shp"),
                dbf=zf.open(f"{stem}.dbf"),
                shx=zf.open(f"{stem}.shx"),
            )
        except Exception as exc:
            print(f"  {stem} unavailable ({exc}); trying next", file=sys.stderr)
            continue

        fields = [f[0] for f in reader.fields[1:]]
        column = fields.index("CONTINENT")
        groups: dict[str, list] = {}
        for record, geometry in zip(reader.records(), reader.shapes()):
            name = record[column]
            # Natural Earth carries a pseudo-continent for open ocean.
            if name == "Seven seas (open ocean)":
                continue
            groups.setdefault(name, []).append(shape(geometry.__geo_interface__))
        print(f"  continents from {stem}: {sorted(groups)}")
        return {name: unary_union(parts) for name, parts in groups.items()}

    raise RuntimeError("could not load any Natural Earth boundaries")


def assign_continents(species: list[dict[str, Any]]) -> None:
    from shapely.geometry import Point
    from shapely.prepared import prep

    polygons = build_continents()
    prepared = {name: prep(geom) for name, geom in polygons.items()}

    failures = 0
    for row in species:
        row["continent"] = None
        if row["lat"] is None or row["lon"] is None:
            continue
        # One unmappable centroid must never cost the whole crawl.
        try:
            point = Point(row["lon"], row["lat"])
            hit = next((n for n, g in prepared.items() if g.contains(point)), None)
            if hit is None:
                # Seabirds and island endemics routinely centre on open water;
                # fall back to whichever landmass is closest.
                hit = min(polygons, key=lambda n: polygons[n].distance(point))
            row["continent"] = hit
        except Exception as exc:
            failures += 1
            if failures <= 5:
                print(f"  could not place {row['sci']} at ({row['lat']}, {row['lon']}): {exc}")
    placed = sum(1 for r in species if r["continent"])
    print(f"  placed {placed}/{len(species)}, {failures} failed")


# ---------------------------------------------------------------- main


def write_output(species: list[dict[str, Any]]) -> None:
    """Most-viewed first, so the pool cutoff is just a slice of the file."""
    species.sort(key=lambda r: r.get("views", 0), reverse=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(species, ensure_ascii=False, indent=None), encoding="utf-8")
    print(f"  wrote {OUTPUT.name} ({OUTPUT.stat().st_size / 1024:.0f} KB)")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--limit", type=int, help="only process the first N species")
    parser.add_argument("--skip-pageviews", action="store_true")
    args = parser.parse_args()

    species = load_species()
    if args.limit:
        species = species[: args.limit]
    print(f"AVONET species: {len(species)}")

    print("Wikidata…")
    wikidata = fetch_wikidata([s["sci"] for s in species])
    for row in species:
        row.update(wikidata.get(row["sci"], {}))
    with_article = [r for r in species if r.get("wiki")]
    print(f"  {len(with_article)} have an English Wikipedia article")

    if not args.skip_pageviews:
        print("Pageviews…")
        views = fetch_pageviews([r["wiki"] for r in with_article])
        for row in species:
            row["views"] = views.get(row.get("wiki", ""), 0)
    else:
        for row in species:
            row["views"] = 0

    # Everything above this line is half an hour of network work. Write it out
    # before doing anything else, so a bug in a later step costs a rerun of
    # that step and not of the entire crawl.
    write_output(species)

    print("Continents…")
    assign_continents(species)
    write_output(species)

    ranked = [r for r in species if r.get("views")]
    print(f"\nwrote {OUTPUT} ({OUTPUT.stat().st_size / 1024:.0f} KB)")
    print(f"{len(ranked)} species with pageviews; top 15:")
    for row in species[:15]:
        print(f"  {row.get('views', 0):>10,}  {row.get('name') or row['sci']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
